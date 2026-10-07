#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
INDISA V9 - pipeline integrado PILOTO CONGELADO

Integra, para PDFs nuevos:
  1) lectura de worklist Excel y localización de páginas de consentimiento;
  2) registro/rectificación por plantilla;
  3) presencia de 5 campos (solo FILLED/EMPTY/UNCERTAIN);
  4) checkbox + firma paciente + firma médico + data_stamp;
  5) fallback geométrico conservador V8.1.6;
  6) anti-duplicado de timbre médico por YOLO + OCR del RUT esperado;
  7) placement seguro usando firma médica predicha;
  8) homografía rectificada -> PDF original;
  9) inserción del PNG correcto sobre COPIAS del PDF;
 10) results.json + audit_log.csv + summary.json + previews.

NO modifica originales.
NO usa anotaciones GT.
NO recalibra thresholds.

Supuesto de layout heredado y validado en V8:
- consentimiento c1: página 1
- consentimiento c2: página 4
- consentimiento c3: página 7, etc.
Es decir: page = 1 + (consent_index - 1) * 3.
Si el layout no coincide o el registro falla, el caso va a revisión.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
import re
import shutil
import sys
import traceback
from collections import defaultdict
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytesseract
from PIL import Image
from ultralytics import YOLO

try:
    import pymupdf as fitz
except Exception:
    import fitz

os.environ.setdefault("WANDB_MODE", "disabled")

VERSION = "V1.1"

FIELDS = [
    "patient_name",
    "patient_rut",
    "procedure_date",
    "doctor_name",
    "doctor_rut",
]

PATIENT_ALT_FIELDS = {"patient_name", "patient_rut"}

# -----------------------------------------------------------------------------
# Configuración congelada: V8.1.1 + policy V8.1.4/V8.1.6
# -----------------------------------------------------------------------------

CONFIG = {
    "templates": {
        "endoscopia_alta": {
            "field_rois": {
                "patient_name": [0.265, 0.157, 0.85, 0.181],
                "procedure_date": [0.185, 0.187, 0.36, 0.211],
                "patient_rut": [0.425, 0.187, 0.69, 0.211],
                "doctor_name": [0.285, 0.816, 0.855, 0.85],
                "doctor_rut": [0.195, 0.852, 0.445, 0.895],
            },
            "auxiliary_rois": {
                "patient_header_zone": [0.29, 0.03, 0.79, 0.125],
                "patient_signature_zone": [0.575, 0.75, 0.875, 0.82],
                "doctor_signature_zone": [0.565, 0.842, 0.875, 0.92],
                "doctor_identity_zone": [0.18, 0.805, 0.875, 0.92],
                "stamp_target_zone": [0.335, 0.825, 0.565, 0.915],
            },
            "checkbox_rois": {
                "Paciente": [0.386, 0.699, 0.414, 0.724],
                "Padre/Madre": [0.535, 0.699, 0.565, 0.724],
                "Apoderado": [0.705, 0.699, 0.735, 0.724],
            },
        },
        "colonoscopia": {
            "field_rois": {
                "patient_name": [0.265, 0.172313, 0.85, 0.196313],
                "procedure_date": [0.185, 0.202313, 0.36, 0.226313],
                "patient_rut": [0.425, 0.202313, 0.69, 0.226313],
                "doctor_name": [0.285, 0.831313, 0.855, 0.865313],
                "doctor_rut": [0.195, 0.867313, 0.445, 0.910313],
            },
            "auxiliary_rois": {
                "patient_header_zone": [0.29, 0.045313, 0.79, 0.140313],
                "patient_signature_zone": [0.575, 0.765313, 0.875, 0.835313],
                "doctor_signature_zone": [0.565, 0.857313, 0.875, 0.935313],
                "doctor_identity_zone": [0.18, 0.820313, 0.875, 0.935313],
                "stamp_target_zone": [0.335, 0.840313, 0.565, 0.930313],
            },
            "checkbox_rois": {
                "Paciente": [0.386, 0.714313, 0.414, 0.739313],
                "Padre/Madre": [0.535, 0.714313, 0.565, 0.739313],
                "Apoderado": [0.705, 0.714313, 0.735, 0.739313],
            },
        },
    }
}

POLICY = {
    "field_presence_accept_confidence": 0.80,

    # V9.0.2: alineados con el pipeline visual congelado que se validó en TEST.
    # Un candidato >=0.03 puede ser rescatado únicamente por la confirmación
    # geométrica conservadora; >=0.20 se acepta como PRESENT.
    "checkbox": {"candidate_floor": 0.03, "high": 0.20},
    "patient_signature": {"candidate_floor": 0.03, "high": 0.20},
    "doctor_signature": {"candidate_floor": 0.03, "high": 0.20},

    # data_stamp a 0.03 replica el piloto no-GT: un candidato inferior débil
    # nunca bloquea automáticamente; requiere confirmación OCR del RUT esperado
    # o deriva a revisión.
    "data_stamp": {"candidate_floor": 0.03, "high": 0.20},
}

GEOM_THRESHOLDS = {
    "patient_signature": 0.021232,
    "doctor_signature": 0.013727,
    "checkbox": 0.023835,
}

GEOM_ROIS = {
    "endoscopia_alta": {
        "patient_signature": (0.54, 0.715, 0.90, 0.805),
        "doctor_signature": (0.47, 0.805, 0.91, 0.920),
        "checkbox_y": (0.674, 0.724),
    },
    "colonoscopia": {
        "patient_signature": (0.54, 0.730, 0.90, 0.815),
        "doctor_signature": (0.47, 0.820, 0.91, 0.930),
        "checkbox_y": (0.688, 0.738),
    },
}

CHECKBOX_X = {
    "Paciente": (0.350, 0.445),
    "Padre/Madre": (0.455, 0.590),
    "Apoderado": (0.615, 0.755),
}

RUT_Y_CENTER = {
    "endoscopia_alta": (0.851852 + 0.894587) / 2.0,
    "colonoscopia": (0.867236 + 0.909972) / 2.0,
}

SAFE_X2 = {"endoscopia_alta": 0.585, "colonoscopia": 0.585}
DESIRED_STAMP_WIDTH_NORM = 0.18
MIN_STAMP_WIDTH_NORM = 0.135
MAX_STAMP_HEIGHT_NORM = 0.072

MIN_H_INLIERS = 15
MIN_H_RATIO = 0.35
MAX_H_MEDIAN_ERROR_PX = 5.0

# -----------------------------------------------------------------------------
# Utilidades
# -----------------------------------------------------------------------------


def now_iso():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def normalize_text(s):
    return (
        (str(s) if s is not None else "")
        .upper()
        .replace("Á", "A").replace("É", "E").replace("Í", "I")
        .replace("Ó", "O").replace("Ú", "U").replace("Ñ", "N")
    )


def norm_rut(x):
    if x is None:
        return ""
    return "".join(c for c in str(x).upper() if c.isdigit() or c == "K")


def attention_id_from_name(name):
    m = re.match(r"(\d+)", Path(name).stem)
    return m.group(1) if m else ""


def safe_imread(path, flags=cv2.IMREAD_COLOR):
    data = np.fromfile(str(path), dtype=np.uint8)
    if data.size == 0:
        return None
    return cv2.imdecode(data, flags)


def safe_imwrite(path, img, quality=95):
    path = Path(path)
    ext = path.suffix.lower() or ".jpg"
    params = []
    if ext in (".jpg", ".jpeg"):
        params = [int(cv2.IMWRITE_JPEG_QUALITY), int(quality)]
    ok, enc = cv2.imencode(ext, img, params)
    if not ok:
        return False
    enc.tofile(str(path))
    return True


def nbox_to_px(box, w, h):
    x1, y1, x2, y2 = map(float, box)
    return [
        max(0, min(w - 1, int(round(x1 * w)))),
        max(0, min(h - 1, int(round(y1 * h)))),
        max(1, min(w, int(round(x2 * w)))),
        max(1, min(h, int(round(y2 * h)))),
    ]


def expand(box, factor):
    x1, y1, x2, y2 = map(float, box)
    rw, rh = x2 - x1, y2 - y1
    return [
        max(0.0, x1 - rw * factor),
        max(0.0, y1 - rh * factor),
        min(1.0, x2 + rw * factor),
        min(1.0, y2 + rh * factor),
    ]


def center_in(xyxy, roi, w, h):
    x1, y1, x2, y2 = map(float, xyxy)
    cx = ((x1 + x2) / 2.0) / w
    cy = ((y1 + y2) / 2.0) / h
    a, b, c, d = roi
    return a <= cx <= c and b <= cy <= d


def bbox_json(xyxy, w, h):
    x1, y1, x2, y2 = map(float, xyxy)
    return {
        "xyxy_px": [round(x1, 1), round(y1, 1), round(x2, 1), round(y2, 1)],
        "xyxy_norm": [
            round(x1 / w, 6), round(y1 / h, 6),
            round(x2 / w, 6), round(y2 / h, 6),
        ],
    }


def box_iou(a, b):
    ax1, ay1, ax2, ay2 = map(float, a)
    bx1, by1, bx2, by2 = map(float, b)
    ix1, iy1 = max(ax1, bx1), max(ay1, by1)
    ix2, iy2 = min(ax2, bx2), min(ay2, by2)
    iw, ih = max(0.0, ix2 - ix1), max(0.0, iy2 - iy1)
    inter = iw * ih
    aa = max(0.0, ax2 - ax1) * max(0.0, ay2 - ay1)
    bb = max(0.0, bx2 - bx1) * max(0.0, by2 - by1)
    den = aa + bb - inter
    return inter / den if den > 0 else 0.0


def dedupe_candidates(cands, threshold=0.75):
    ordered = sorted(
        cands or [], key=lambda x: float(x.get("confidence", 0) or 0), reverse=True
    )
    keep = []
    for c in ordered:
        box = c.get("bbox", {}).get("xyxy_px")
        if not box:
            keep.append(c)
            continue
        duplicate = False
        for k in keep:
            kbox = k.get("bbox", {}).get("xyxy_px")
            if kbox and box_iou(box, kbox) >= threshold:
                duplicate = True
                break
        if not duplicate:
            keep.append(c)
    return keep


# -----------------------------------------------------------------------------
# Excel / timbres / PDFs
# -----------------------------------------------------------------------------


def load_worklist(xlsx_path):
    """
    Carga el Excel de atenciones de forma robusta.

    Fix V9.0.1:
    - no asume que todas las filas tienen el mismo largo;
    - no indexa directamente una tupla corta;
    - busca automáticamente la hoja que contiene las columnas obligatorias;
    - tolera filas vacías o parcialmente vacías.
    """
    wb = openpyxl.load_workbook(
        xlsx_path,
        read_only=True,
        data_only=True,
    )

    required = ["id", "procedimiento", "rut médico"]

    ws_selected = None
    headers = None
    idx = None
    header_row_number = None

    # Buscar una hoja válida y, por seguridad, revisar las primeras 20 filas
    # por si el encabezado no estuviera exactamente en la fila 1.
    for ws in wb.worksheets:
        for row_number, row in enumerate(
            ws.iter_rows(min_row=1, max_row=20, values_only=True),
            start=1,
        ):
            if not row:
                continue

            candidate_headers = [
                str(x).strip() if x is not None else ""
                for x in row
            ]
            candidate_idx = {
                h.lower(): i
                for i, h in enumerate(candidate_headers)
                if h
            }

            if all(col in candidate_idx for col in required):
                ws_selected = ws
                headers = candidate_headers
                idx = candidate_idx
                header_row_number = row_number
                break

        if ws_selected is not None:
            break

    if ws_selected is None:
        sheet_info = ", ".join(wb.sheetnames)
        raise ValueError(
            "No encontré una hoja con las columnas obligatorias "
            f"{required}. Hojas disponibles: {sheet_info}"
        )

    def cell_value(row, col_name):
        col_idx = idx.get(col_name)
        if col_idx is None or col_idx >= len(row):
            return None
        return row[col_idx]

    out = defaultdict(list)

    for row in ws_selected.iter_rows(
        min_row=header_row_number + 1,
        values_only=True,
    ):
        if not row or all(v is None for v in row):
            continue

        rid = cell_value(row, "id")
        if rid is None:
            continue

        if isinstance(rid, float) and rid.is_integer():
            aid = str(int(rid))
        else:
            aid = str(rid).strip()

        if not aid:
            continue

        def val(col):
            v = cell_value(row, col)
            return "" if v is None else str(v).strip()

        out[aid].append({
            "attention_id": aid,
            "procedure": val("procedimiento"),
            "doctor_rut": norm_rut(val("rut médico")),
            "doctor_name": val("nombre médico"),
            "patient_rut": norm_rut(val("rut paciente")),
            "patient_name": val("nombre paciente"),
        })

    if not out:
        raise RuntimeError(
            f"No se cargaron atenciones desde la hoja "
            f"'{ws_selected.title}'."
        )

    print(
        f"Excel atenciones: hoja='{ws_selected.title}' "
        f"| IDs cargados={len(out)}"
    )

    return out

def load_stamp_files(folder):
    out = {}
    for p in Path(folder).glob("*.png"):
        m = re.match(r"([^_]+)_", p.name)
        if not m:
            continue
        rut = norm_rut(m.group(1))
        if rut:
            out[rut] = p
    return out


def collect_input_pdfs(input_path):
    p = Path(input_path)
    if p.is_file():
        return [p] if p.suffix.lower() == ".pdf" else []
    pdfs = []
    for x in p.rglob("*.pdf"):
        low = str(x).lower()
        if "plantillas base" in low or "consentimiento_base" in low:
            continue
        pdfs.append(x)
    return sorted(pdfs)


def render_pdf_page_bgr(pdf_path, page_number_1idx, dpi=150):
    doc = fitz.open(str(pdf_path))
    try:
        page = doc[page_number_1idx - 1]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width, pix.n
        )
        if pix.n == 4:
            return cv2.cvtColor(arr, cv2.COLOR_RGBA2BGR)
        return cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)
    finally:
        doc.close()


# -----------------------------------------------------------------------------
# Rectificación por plantilla
# -----------------------------------------------------------------------------


def orb_register(template_bgr, scan_bgr):
    def prep(im, width=800):
        gray = cv2.cvtColor(im, cv2.COLOR_BGR2GRAY)
        scale = width / gray.shape[1]
        return cv2.resize(gray, None, fx=scale, fy=scale), scale

    tpl, st = prep(template_bgr)
    scn, ss = prep(scan_bgr)

    orb = cv2.ORB_create(nfeatures=2000, fastThreshold=12)
    kp1, d1 = orb.detectAndCompute(tpl, None)
    kp2, d2 = orb.detectAndCompute(scn, None)
    if d1 is None or d2 is None:
        raise RuntimeError("ORB sin descriptores suficientes")

    bf = cv2.BFMatcher(cv2.NORM_HAMMING)
    knn = bf.knnMatch(d1, d2, k=2)
    good = []
    for pair in knn:
        if len(pair) < 2:
            continue
        m, n = pair
        if m.distance < 0.76 * n.distance:
            good.append(m)

    if len(good) < 10:
        raise RuntimeError(f"Registro ORB débil: {len(good)} matches")

    src_small = np.float32([kp1[m.queryIdx].pt for m in good]).reshape(-1, 1, 2)
    dst_small = np.float32([kp2[m.trainIdx].pt for m in good]).reshape(-1, 1, 2)

    Hs, mask = cv2.findHomography(src_small, dst_small, cv2.RANSAC, 4.0, maxIters=2000)
    if Hs is None:
        raise RuntimeError("No se pudo estimar homografía ORB")

    Ssrc = np.array([[st, 0, 0], [0, st, 0], [0, 0, 1]], dtype=np.float64)
    Sdst_inv = np.array([[1 / ss, 0, 0], [0, 1 / ss, 0], [0, 0, 1]], dtype=np.float64)
    H = Sdst_inv @ Hs @ Ssrc
    inliers = int(mask.sum()) if mask is not None else 0
    return H, len(good), inliers


def rectify_scan(scan_bgr, H_template_to_scan, template_shape):
    th, tw = template_shape[:2]
    H_scan_to_template = np.linalg.inv(H_template_to_scan)
    return cv2.warpPerspective(
        scan_bgr,
        H_scan_to_template,
        (tw, th),
        flags=cv2.INTER_LINEAR,
        borderMode=cv2.BORDER_CONSTANT,
        borderValue=(255, 255, 255),
    )


def infer_template_type(procedure):
    return "colonoscopia" if "COLON" in normalize_text(procedure) else "endoscopia_alta"


# -----------------------------------------------------------------------------
# Modelos: campos + visual
# -----------------------------------------------------------------------------


def model_imgsz(model, default=960):
    try:
        v = model.ckpt.get("train_args", {}).get("imgsz", default)
        if isinstance(v, (list, tuple)):
            return int(max(v))
        return int(v)
    except Exception:
        return default


def classify_field(model, crop):
    r = model.predict(source=crop, verbose=False)[0]
    idx = int(r.probs.top1)
    conf = float(r.probs.top1conf.item())
    label = str(model.names[idx]).lower()

    if conf < POLICY["field_presence_accept_confidence"]:
        return {
            "status": "UNCERTAIN",
            "predicted_class": label,
            "confidence": round(conf, 6),
            "manual_review_required": True,
            "review_reason": "low_confidence",
            "source_model": "field_presence_v811",
        }
    if label == "empty":
        return {
            "status": "EMPTY",
            "predicted_class": label,
            "confidence": round(conf, 6),
            "manual_review_required": True,
            "review_reason": "missing_required_field",
            "source_model": "field_presence_v811",
        }
    return {
        "status": "PRESENT",
        "predicted_class": label,
        "confidence": round(conf, 6),
        "manual_review_required": False,
        "review_reason": None,
        "source_model": "field_presence_v811",
    }


def predict_candidates(model, img, conf, device):
    r = model.predict(
        source=img,
        conf=conf,
        iou=0.7,
        imgsz=model_imgsz(model),
        device=device,
        verbose=False,
    )[0]
    out = []
    if r.boxes is None:
        return out
    for b in r.boxes:
        out.append({
            "class_id": int(b.cls.item()),
            "confidence": float(b.conf.item()),
            "xyxy": b.xyxy[0].detach().cpu().numpy(),
        })
    return out


def summarize_one(cands, cls_id, floor, high, zone, w, h, source):
    vals = []
    for c in cands:
        if c["class_id"] != cls_id or c["confidence"] < floor:
            continue
        if zone is not None and not center_in(c["xyxy"], zone, w, h):
            continue
        vals.append({
            "confidence": round(c["confidence"], 6),
            "bbox": bbox_json(c["xyxy"], w, h),
        })
    vals.sort(key=lambda x: x["confidence"], reverse=True)

    if not vals:
        return {
            "status": "NOT_DETECTED",
            "detected": False,
            "confidence": None,
            "manual_review_required": True,
            "review_reason": "not_detected",
            "source_model": source,
            "candidates": [],
        }

    top = vals[0]
    if top["confidence"] >= high:
        return {
            "status": "PRESENT",
            "detected": True,
            "confidence": top["confidence"],
            "manual_review_required": False,
            "review_reason": None,
            "source_model": source,
            "candidates": vals,
        }

    return {
        "status": "PRESENT_LOW_CONFIDENCE",
        "detected": True,
        "confidence": top["confidence"],
        "manual_review_required": True,
        "review_reason": "low_confidence",
        "source_model": source,
        "candidates": vals,
    }


def checkbox_option(xyxy, rois, w, h):
    for name, roi in rois.items():
        if center_in(xyxy, expand(roi, 0.35), w, h):
            return name
    return None


def summarize_checkbox(cands, tc, w, h):
    p = POLICY["checkbox"]
    vals = []
    for c in cands:
        if c["class_id"] != 0 or c["confidence"] < p["candidate_floor"]:
            continue
        opt = checkbox_option(c["xyxy"], tc["checkbox_rois"], w, h)
        if opt is None:
            continue
        vals.append({
            "option": opt,
            "confidence": round(c["confidence"], 6),
            "bbox": bbox_json(c["xyxy"], w, h),
        })

    vals.sort(key=lambda x: x["confidence"], reverse=True)
    if not vals:
        return {
            "status": "NOT_DETECTED",
            "selected_option": None,
            "confidence": None,
            "manual_review_required": True,
            "review_reason": "not_detected",
            "source_model": "visual_main_v813",
            "candidates": [],
        }

    top = vals[0]
    opts = {x["option"] for x in vals}
    if len(opts) > 1:
        return {
            "status": "AMBIGUOUS",
            "selected_option": top["option"],
            "confidence": top["confidence"],
            "manual_review_required": True,
            "review_reason": "ambiguous_multiple_candidates",
            "source_model": "visual_main_v813",
            "candidates": vals,
        }

    low = top["confidence"] < p["high"]
    return {
        "status": "PRESENT_LOW_CONFIDENCE" if low else "PRESENT",
        "selected_option": top["option"],
        "confidence": top["confidence"],
        "manual_review_required": low,
        "review_reason": "low_confidence" if low else None,
        "source_model": "visual_main_v813",
        "candidates": vals,
    }


def summarize_stamps(cands, w, h):
    p = POLICY["data_stamp"]
    vals = [
        {
            "confidence": round(c["confidence"], 6),
            "bbox": bbox_json(c["xyxy"], w, h),
        }
        for c in cands
        if c["class_id"] == 3 and c["confidence"] >= p["candidate_floor"]
    ]
    vals.sort(key=lambda x: x["confidence"], reverse=True)
    vals = dedupe_candidates(vals, 0.75)

    if not vals:
        return {
            "status": "NONE_DETECTED",
            "count": 0,
            "confidence": None,
            "manual_review_required": False,
            "review_reason": None,
            "source_model": "visual_main_v813",
            "candidates": [],
        }

    high = [x for x in vals if x["confidence"] >= p["high"]]
    return {
        "status": "DETECTED" if high else "DETECTED_LOW_CONFIDENCE",
        "count": len(vals),
        "confidence": vals[0]["confidence"],
        "manual_review_required": False,
        "review_reason": None,
        "source_model": "visual_main_v813",
        "candidates": vals,
    }


# -----------------------------------------------------------------------------
# Presence-only + fallback geométrico V8.1.6
# -----------------------------------------------------------------------------


def top_patient_stamp_candidates(doc, y_limit=0.20):
    ds = doc.get("visual", {}).get("data_stamp", {})
    out = []
    for c in dedupe_candidates(ds.get("candidates", []), 0.75):
        n = c.get("bbox", {}).get("xyxy_norm")
        if not n or len(n) != 4:
            continue
        cy = (float(n[1]) + float(n[3])) / 2
        if cy <= y_limit:
            out.append(c)
    out.sort(key=lambda x: float(x.get("confidence", 0) or 0), reverse=True)
    return out


def simplify_field_presence(name, src, top_stamp):
    predicted = str(src.get("predicted_class", "")).lower()
    confidence = src.get("confidence")
    status = src.get("status")

    if predicted == "present" or status == "PRESENT":
        return {
            "status": "FILLED",
            "filled": True,
            "confidence": confidence,
            "source": "primary_field_roi",
            "manual_review_required": False,
            "review_reason": None,
            "source_model": src.get("source_model"),
            "roi_px": src.get("roi_px"),
        }

    if name in PATIENT_ALT_FIELDS and top_stamp:
        best = top_stamp[0]
        return {
            "status": "FILLED",
            "filled": True,
            "confidence": best.get("confidence"),
            "source": "top_patient_data_stamp",
            "primary_field_status": status,
            "manual_review_required": False,
            "review_reason": None,
            "source_model": "visual_main_v813",
            "roi_px": src.get("roi_px"),
        }

    if predicted == "empty" or status == "EMPTY":
        return {
            "status": "EMPTY",
            "filled": False,
            "confidence": confidence,
            "source": "primary_field_roi",
            "manual_review_required": True,
            "review_reason": "missing_required_field",
            "source_model": src.get("source_model"),
            "roi_px": src.get("roi_px"),
        }

    return {
        "status": "UNCERTAIN",
        "filled": None,
        "confidence": confidence,
        "source": "primary_field_roi",
        "manual_review_required": True,
        "review_reason": "low_confidence",
        "source_model": src.get("source_model"),
        "roi_px": src.get("roi_px"),
    }


def align_template_ecc(template_bgr, real_bgr):
    r = cv2.cvtColor(real_bgr, cv2.COLOR_BGR2GRAY)
    t = cv2.cvtColor(template_bgr, cv2.COLOR_BGR2GRAY)
    r = cv2.GaussianBlur(r, (3, 3), 0).astype(np.float32) / 255.0
    t = cv2.GaussianBlur(t, (3, 3), 0).astype(np.float32) / 255.0
    warp = np.eye(2, 3, dtype=np.float32)
    criteria = (cv2.TERM_CRITERIA_EPS | cv2.TERM_CRITERIA_COUNT, 150, 1e-6)
    try:
        cv2.findTransformECC(r, t, warp, cv2.MOTION_AFFINE, criteria, None, 1)
        h, w = real_bgr.shape[:2]
        return cv2.warpAffine(
            template_bgr,
            warp,
            (w, h),
            flags=cv2.INTER_LINEAR + cv2.WARP_INVERSE_MAP,
            borderMode=cv2.BORDER_CONSTANT,
            borderValue=(255, 255, 255),
        )
    except cv2.error:
        return template_bgr


def diff_features(real_crop, tmpl_crop):
    rg = cv2.cvtColor(real_crop, cv2.COLOR_BGR2GRAY)
    tg = cv2.cvtColor(tmpl_crop, cv2.COLOR_BGR2GRAY)

    shift = float(np.median(rg)) - float(np.median(tg))
    ta = np.clip(tg.astype(np.float32) + shift, 0, 255).astype(np.uint8)
    delta = np.clip(ta.astype(np.int16) - rg.astype(np.int16), 0, 255).astype(np.uint8)
    raw = (delta >= 26).astype(np.uint8)

    template_ink = (tg < 220).astype(np.uint8)
    template_ink = cv2.dilate(
        template_ink,
        cv2.getStructuringElement(cv2.MORPH_RECT, (3, 3)),
        iterations=1,
    )

    clean = raw.copy()
    clean[template_ink > 0] = 0

    n, labels, stats, _ = cv2.connectedComponentsWithStats(clean, 8)
    filtered = np.zeros_like(clean)
    for i in range(1, n):
        area = int(stats[i, cv2.CC_STAT_AREA])
        ww = int(stats[i, cv2.CC_STAT_WIDTH])
        hh = int(stats[i, cv2.CC_STAT_HEIGHT])
        if area >= 3 and (ww >= 2 or hh >= 2):
            filtered[labels == i] = 1

    return {"raw_ratio": float(raw.mean()), "clean_ratio": float(filtered.mean())}


def geometric_scores(real, template_aligned, template_type):
    h, w = real.shape[:2]
    cfg = GEOM_ROIS[template_type]
    scores = {}

    for comp in ("patient_signature", "doctor_signature"):
        b = nbox_to_px(cfg[comp], w, h)
        f = diff_features(
            real[b[1]:b[3], b[0]:b[2]],
            template_aligned[b[1]:b[3], b[0]:b[2]],
        )
        scores[comp] = f["clean_ratio"]

    y1, y2 = cfg["checkbox_y"]
    cb = {}
    for option, (x1, x2) in CHECKBOX_X.items():
        b = nbox_to_px((x1, y1, x2, y2), w, h)
        f = diff_features(
            real[b[1]:b[3], b[0]:b[2]],
            template_aligned[b[1]:b[3], b[0]:b[2]],
        )
        cb[option] = f["raw_ratio"]

    best = max(cb, key=cb.get)
    ordered = sorted(cb.values(), reverse=True)
    scores["checkbox"] = {
        "scores_by_option": cb,
        "best_option": best,
        "presence_score": cb[best],
        "option_margin": ordered[0] - ordered[1] if len(ordered) >= 2 else 0.0,
    }
    return scores


def maybe_confirm_signature(comp, score, threshold):
    if not comp or comp.get("status") != "PRESENT_LOW_CONFIDENCE":
        return
    if not comp.get("detected", False):
        return

    confirmed = bool(score >= threshold)
    comp["geometric_confirmation"] = {
        "score": round(float(score), 6),
        "threshold": threshold,
        "confirmed": confirmed,
        "rule": "confirm_existing_low_confidence_yolo_candidate_only",
    }
    if confirmed:
        comp["status"] = "PRESENT_GEOMETRIC_CONFIRMED"
        comp["manual_review_required"] = False
        comp["review_reason"] = None


def maybe_confirm_checkbox(comp, cb):
    if not comp or comp.get("status") != "PRESENT_LOW_CONFIDENCE":
        return

    yolo_option = comp.get("selected_option")
    score = float(cb["presence_score"])
    geom_option = cb["best_option"]
    confirmed = (
        yolo_option is not None
        and geom_option == yolo_option
        and score >= GEOM_THRESHOLDS["checkbox"]
    )

    comp["geometric_confirmation"] = {
        "presence_score": round(score, 6),
        "threshold": GEOM_THRESHOLDS["checkbox"],
        "geometry_best_option": geom_option,
        "yolo_selected_option": yolo_option,
        "option_agreement": geom_option == yolo_option,
        "option_margin": round(float(cb["option_margin"]), 6),
        "confirmed": confirmed,
        "rule": "confirm_existing_low_confidence_yolo_candidate_only",
    }
    if confirmed:
        comp["status"] = "PRESENT_GEOMETRIC_CONFIRMED"
        comp["manual_review_required"] = False
        comp["review_reason"] = None


# -----------------------------------------------------------------------------
# Anti-duplicado del timbre
# -----------------------------------------------------------------------------


def lower_stamp_candidates(doc, y_limit=0.30):
    cands = dedupe_candidates(
        doc.get("visual", {}).get("data_stamp", {}).get("candidates", []), 0.75
    )
    out = []
    for c in cands:
        n = c.get("bbox", {}).get("xyxy_norm")
        if not n:
            continue
        cy = (float(n[1]) + float(n[3])) / 2.0
        if cy > y_limit:
            out.append(c)
    return out


def boxes_close(a, b, x_gap=0.04, y_gap=0.035):
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    dx = max(0.0, max(ax1, bx1) - min(ax2, bx2))
    dy = max(0.0, max(ay1, by1) - min(ay2, by2))
    return box_iou(a, b) > 0 or (dx <= x_gap and dy <= y_gap)


def cluster_stamp_candidates(cands):
    boxes = [tuple(c["bbox"]["xyxy_norm"]) for c in cands]
    used = [False] * len(boxes)
    clusters = []

    for i in range(len(boxes)):
        if used[i]:
            continue
        used[i] = True
        group = [i]
        changed = True
        while changed:
            changed = False
            for j in range(len(boxes)):
                if used[j]:
                    continue
                if any(boxes_close(boxes[j], boxes[k]) for k in group):
                    used[j] = True
                    group.append(j)
                    changed = True
        clusters.append([cands[k] for k in group])
    return clusters


def union_crop_norm(img, cluster, pad=0.025):
    h, w = img.shape[:2]
    boxes = [c["bbox"]["xyxy_norm"] for c in cluster]
    x1 = max(0.0, min(b[0] for b in boxes) - pad)
    y1 = max(0.0, min(b[1] for b in boxes) - pad)
    x2 = min(1.0, max(b[2] for b in boxes) + pad)
    y2 = min(1.0, max(b[3] for b in boxes) + pad)
    X1, Y1, X2, Y2 = nbox_to_px((x1, y1, x2, y2), w, h)
    return img[Y1:Y2, X1:X2].copy(), [x1, y1, x2, y2]


def ocr_expected_rut(crop, expected_rut):
    if crop is None or crop.size == 0 or not expected_rut:
        return {"confirmed": False, "hit_count": 0, "variants_hit": []}

    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    up = cv2.resize(gray, None, fx=3.0, fy=3.0, interpolation=cv2.INTER_CUBIC)
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(up)
    _, otsu = cv2.threshold(clahe, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    adaptive = cv2.adaptiveThreshold(
        clahe, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY, 35, 11,
    )
    sharp = cv2.addWeighted(
        clahe, 1.8, cv2.GaussianBlur(clahe, (0, 0), 1.1), -0.8, 0
    )

    variants = {
        "gray_up": up,
        "clahe": clahe,
        "otsu": otsu,
        "adaptive": adaptive,
        "sharp": sharp,
    }

    hits = []
    for vname, im in variants.items():
        for psm in (6, 11, 12):
            txt = pytesseract.image_to_string(
                im, lang="eng", config=f"--oem 3 --psm {psm}"
            )
            if expected_rut in norm_rut(txt):
                hits.append((vname, "general", psm))

        for psm in (6, 7, 11, 12, 13):
            txt = pytesseract.image_to_string(
                im,
                lang="eng",
                config=(
                    f"--oem 3 --psm {psm} "
                    "-c tessedit_char_whitelist=0123456789Kk.-"
                ),
            )
            if expected_rut in norm_rut(txt):
                hits.append((vname, "rut_whitelist", psm))

    variants_hit = sorted({x[0] for x in hits})
    confirmed = len(hits) >= 2 and len(variants_hit) >= 2
    return {
        "confirmed": confirmed,
        "hit_count": len(hits),
        "variants_hit": variants_hit,
    }


def inspect_existing_doctor_stamp(rectified, doc, expected_rut):
    cands = lower_stamp_candidates(doc)
    if not cands:
        return {
            "status": "NONE",
            "candidate_count": 0,
            "cluster_count": 0,
            "confirmed": False,
            "clusters": [],
        }

    clusters = cluster_stamp_candidates(cands)
    details = []
    any_confirmed = False

    for cluster in clusters:
        crop, ub = union_crop_norm(rectified, cluster)
        try:
            ocr = ocr_expected_rut(crop, expected_rut)
        except Exception as exc:
            ocr = {
                "confirmed": False,
                "hit_count": 0,
                "variants_hit": [],
                "error": f"{type(exc).__name__}: {exc}",
            }
        any_confirmed = any_confirmed or bool(ocr.get("confirmed"))
        details.append({
            "candidate_boxes": len(cluster),
            "union_box_norm": [round(float(x), 6) for x in ub],
            **ocr,
        })

    return {
        "status": "CONFIRMED_EXPECTED_DOCTOR_STAMP" if any_confirmed else "INCONCLUSIVE_CANDIDATE",
        "candidate_count": len(cands),
        "cluster_count": len(clusters),
        "confirmed": any_confirmed,
        "clusters": details,
    }


# -----------------------------------------------------------------------------
# Timbre / placement
# -----------------------------------------------------------------------------


def normalize_stamp_gray(path):
    im = safe_imread(path, cv2.IMREAD_GRAYSCALE)
    if im is None or im.size == 0:
        return None

    border = np.concatenate([im[0, :], im[-1, :], im[:, 0], im[:, -1]])
    if float(np.median(border)) < 128:
        im = 255 - im

    lo, hi = np.percentile(im, [1, 99.5])
    if hi > lo:
        im = np.clip(
            (im.astype(np.float32) - lo) * 255.0 / (hi - lo), 0, 255
        ).astype(np.uint8)

    mask = im < 248
    ys, xs = np.where(mask)
    if len(xs):
        pad = 4
        x1 = max(0, int(xs.min()) - pad)
        y1 = max(0, int(ys.min()) - pad)
        x2 = min(im.shape[1], int(xs.max()) + 1 + pad)
        y2 = min(im.shape[0], int(ys.max()) + 1 + pad)
        im = im[y1:y2, x1:x2]
    return im


def best_signature_bbox_norm(doc):
    comp = doc.get("visual", {}).get("doctor_signature", {})
    if comp.get("status") not in {"PRESENT", "PRESENT_GEOMETRIC_CONFIRMED"}:
        return None
    cands = comp.get("candidates", []) or []
    if not cands:
        return None
    best = max(cands, key=lambda x: float(x.get("confidence", 0) or 0))
    box = best.get("bbox", {}).get("xyxy_norm")
    return tuple(map(float, box)) if box and len(box) == 4 else None


def propose_stamp_target(page_shape, stamp_gray, template_type, signature_box=None, allow_fixed=False):
    h, w = page_shape[:2]
    if stamp_gray is None or stamp_gray.size == 0:
        return None, "stamp_image_unreadable", None

    sh, sw = stamp_gray.shape[:2]
    width = DESIRED_STAMP_WIDTH_NORM
    height = (width * w) * (sh / sw) / h

    if height > MAX_STAMP_HEIGHT_NORM:
        scale = MAX_STAMP_HEIGHT_NORM / height
        width *= scale
        height *= scale

    if signature_box is not None:
        sx1 = float(signature_box[0])
        x2 = min(SAFE_X2[template_type], sx1 - 0.015)
        mode = "signature_aware"
    elif allow_fixed:
        x2 = SAFE_X2[template_type]
        mode = "fixed_safe_zone_no_signature"
    else:
        return None, "doctor_signature_not_safely_located", None

    x1 = x2 - width
    if x1 < 0.30:
        d = 0.30 - x1
        x1 += d
        x2 += d

    if x2 > SAFE_X2[template_type]:
        d = x2 - SAFE_X2[template_type]
        x1 -= d
        x2 -= d

    if (x2 - x1) < MIN_STAMP_WIDTH_NORM:
        return None, "insufficient_horizontal_space", mode

    cy = RUT_Y_CENTER[template_type]
    y1 = cy - height / 2
    y2 = cy + height / 2

    if y2 > 0.935:
        d = y2 - 0.935
        y1 -= d
        y2 -= d

    if y1 < 0.79:
        d = 0.79 - y1
        y1 += d
        y2 += d

    return (x1, y1, x2, y2), "ok", mode


# -----------------------------------------------------------------------------
# Homografía rectificada -> PDF original
# -----------------------------------------------------------------------------


def resize_for_features(gray, max_side=1800):
    h, w = gray.shape[:2]
    s = min(1.0, max_side / max(h, w))
    if s < 1.0:
        return cv2.resize(
            gray, (int(round(w * s)), int(round(h * s))), interpolation=cv2.INTER_AREA
        ), s
    return gray, s


def estimate_pdf_homography(rectified, raw):
    g1 = cv2.cvtColor(rectified, cv2.COLOR_BGR2GRAY)
    g2 = cv2.cvtColor(raw, cv2.COLOR_BGR2GRAY)
    f1, s1 = resize_for_features(g1)
    f2, s2 = resize_for_features(g2)

    sift = cv2.SIFT_create(nfeatures=5000, contrastThreshold=0.015, edgeThreshold=12)
    k1, d1 = sift.detectAndCompute(f1, None)
    k2, d2 = sift.detectAndCompute(f2, None)
    if d1 is None or d2 is None or len(k1) < 12 or len(k2) < 12:
        return None, {"reason": "insufficient_features"}

    knn = cv2.BFMatcher(cv2.NORM_L2).knnMatch(d1, d2, k=2)
    good = []
    for pair in knn:
        if len(pair) < 2:
            continue
        m, n = pair
        if m.distance < 0.72 * n.distance:
            good.append(m)

    if len(good) < 10:
        return None, {"reason": "insufficient_good_matches", "good_matches": len(good)}

    src = np.float32([k1[m.queryIdx].pt for m in good])
    dst = np.float32([k2[m.trainIdx].pt for m in good])
    src[:, 0] /= s1
    src[:, 1] /= s1
    dst[:, 0] /= s2
    dst[:, 1] /= s2

    H, mask = cv2.findHomography(src, dst, cv2.RANSAC, 4.0)
    if H is None or mask is None:
        return None, {"reason": "homography_failed"}

    ins = mask.ravel().astype(bool)
    nin = int(ins.sum())
    ratio = nin / len(good)
    if nin:
        pred = cv2.perspectiveTransform(src[ins].reshape(-1, 1, 2), H).reshape(-1, 2)
        err = float(np.median(np.linalg.norm(pred - dst[ins], axis=1)))
    else:
        err = 999.0

    return H, {
        "reason": "ok",
        "good_matches": len(good),
        "inliers": nin,
        "inlier_ratio": ratio,
        "median_reprojection_error_px": err,
    }


def target_quad(box, shape):
    h, w = shape[:2]
    x1, y1, x2, y2 = box
    return np.float32([
        [x1 * w, y1 * h],
        [x2 * w, y1 * h],
        [x2 * w, y2 * h],
        [x1 * w, y2 * h],
    ]).reshape(-1, 1, 2)


def stamp_rgba_from_gray(gray):
    alpha = np.clip(
        ((255 - gray).astype(np.float32) - 10.0) / 245.0 * 255.0, 0, 255
    ).astype(np.uint8)
    rgba = np.zeros((gray.shape[0], gray.shape[1], 4), dtype=np.uint8)
    rgba[..., 3] = alpha
    return rgba


def build_page_overlay_png(raw_shape, stamp_gray, quad):
    h, w = raw_shape[:2]
    rgba = stamp_rgba_from_gray(stamp_gray)
    sh, sw = rgba.shape[:2]

    src = np.float32([[0, 0], [sw - 1, 0], [sw - 1, sh - 1], [0, sh - 1]])
    M = cv2.getPerspectiveTransform(src, quad.astype(np.float32))

    warped_alpha = cv2.warpPerspective(
        rgba[..., 3],
        M,
        (w, h),
        flags=cv2.INTER_LINEAR,
        borderMode=cv2.BORDER_CONSTANT,
        borderValue=0,
    )

    overlay = np.zeros((h, w, 4), dtype=np.uint8)
    overlay[..., 3] = warped_alpha

    img = Image.fromarray(cv2.cvtColor(overlay, cv2.COLOR_BGRA2RGBA), mode="RGBA")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# -----------------------------------------------------------------------------
# Análisis completo de una página rectificada
# -----------------------------------------------------------------------------


def analyze_rectified(
    rectified,
    template_bgr,
    template_type,
    field_model,
    visual_main,
    doctor_model,
    device,
):
    h, w = rectified.shape[:2]
    tc = CONFIG["templates"][template_type]

    result = {
        "schema_version": VERSION,
        "field_requirement_mode": "presence_only",
        "fields": {},
        "visual": {},
    }

    # Campos estructurados.
    raw_fields = {}
    for field in FIELDS:
        x1, y1, x2, y2 = nbox_to_px(tc["field_rois"][field], w, h)
        crop = rectified[y1:y2, x1:x2]
        f = classify_field(field_model, crop)
        f["roi_px"] = [x1, y1, x2, y2]
        raw_fields[field] = f

    # YOLO principal una sola vez.
    main_floor = min(
        POLICY["checkbox"]["candidate_floor"],
        POLICY["patient_signature"]["candidate_floor"],
        POLICY["data_stamp"]["candidate_floor"],
    )
    main_cands = predict_candidates(visual_main, rectified, main_floor, device)

    result["visual"]["signer_selected_checkbox"] = summarize_checkbox(
        main_cands, tc, w, h
    )
    result["visual"]["patient_signature"] = summarize_one(
        main_cands,
        1,
        POLICY["patient_signature"]["candidate_floor"],
        POLICY["patient_signature"]["high"],
        expand(tc["auxiliary_rois"]["patient_signature_zone"], 0.35),
        w,
        h,
        "visual_main_v813",
    )
    result["visual"]["data_stamp"] = summarize_stamps(main_cands, w, h)

    doctor_cands = predict_candidates(
        doctor_model, rectified, POLICY["doctor_signature"]["candidate_floor"], device
    )
    result["visual"]["doctor_signature"] = summarize_one(
        doctor_cands,
        2,
        POLICY["doctor_signature"]["candidate_floor"],
        POLICY["doctor_signature"]["high"],
        expand(tc["auxiliary_rois"]["doctor_signature_zone"], 0.50),
        w,
        h,
        "doctor_signature_fallback_v811",
    )

    # Presence-only con etiqueta superior alternativa para paciente.
    top_stamp = top_patient_stamp_candidates(result)
    for name in FIELDS:
        result["fields"][name] = simplify_field_presence(
            name, raw_fields[name], top_stamp
        )

    # Fallback geométrico conservador.
    tmpl_aligned = align_template_ecc(template_bgr, rectified)
    gs = geometric_scores(rectified, tmpl_aligned, template_type)
    maybe_confirm_signature(
        result["visual"]["patient_signature"],
        gs["patient_signature"],
        GEOM_THRESHOLDS["patient_signature"],
    )
    maybe_confirm_signature(
        result["visual"]["doctor_signature"],
        gs["doctor_signature"],
        GEOM_THRESHOLDS["doctor_signature"],
    )
    maybe_confirm_checkbox(result["visual"]["signer_selected_checkbox"], gs["checkbox"])

    result["visual_fallback"] = {
        "mode": "conservative_confirmation_only",
        "patient_signature_score": round(gs["patient_signature"], 6),
        "doctor_signature_score": round(gs["doctor_signature"], 6),
        "checkbox": {
            "best_option": gs["checkbox"]["best_option"],
            "presence_score": round(gs["checkbox"]["presence_score"], 6),
            "option_margin": round(gs["checkbox"]["option_margin"], 6),
        },
        "thresholds_source": "TRAIN_only_calibration_v2_frozen",
    }

    return result


def recalc_manual_review(result):
    items = []

    for name, field in result.get("fields", {}).items():
        if field.get("manual_review_required"):
            items.append({"component": name, "reason": field.get("review_reason")})

    for name, comp in result.get("visual", {}).items():
        if name == "data_stamp":
            continue
        if comp.get("manual_review_required"):
            items.append({"component": name, "reason": comp.get("review_reason")})

    stamp = result.get("stamp_stage", {})
    if stamp.get("manual_review_required"):
        stamp_reason = stamp.get("review_reason")

        # V9.0.3: evitar duplicar una causa ya reportada por doctor_signature.
        # Si stamp_stage está en revisión únicamente porque no se pudo ubicar
        # el timbre al faltar una firma médica segura, la causa raíz ya está
        # representada por doctor_signature.
        doctor_signature_already_reported = any(
            x.get("component") == "doctor_signature"
            for x in items
        )
        duplicate_signature_dependency = (
            doctor_signature_already_reported
            and stamp_reason == "doctor_signature_not_safely_located"
        )

        if not duplicate_signature_dependency:
            items.append({
                "component": "stamp_stage",
                "reason": stamp_reason,
            })

    # Deduplicar pares component/reason.
    seen = set()
    uniq = []
    for x in items:
        key = (x.get("component"), x.get("reason"))
        if key in seen:
            continue
        seen.add(key)
        uniq.append(x)

    result["manual_review_required"] = bool(uniq)
    result["manual_review_items"] = uniq


# -----------------------------------------------------------------------------
# Previews y logs
# -----------------------------------------------------------------------------


def render_doc_page_rgb(doc, page_idx, dpi=160):
    page = doc[page_idx]
    pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72.0, dpi / 72.0), alpha=False)
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
        pix.height, pix.width, pix.n
    )
    if pix.n == 4:
        return cv2.cvtColor(arr, cv2.COLOR_RGBA2RGB)
    return arr[..., :3].copy()


def save_preview_from_pdf(pdf_path, page_idx, out_path):
    doc = fitz.open(str(pdf_path))
    try:
        rgb = render_doc_page_rgb(doc, page_idx, dpi=160)
    finally:
        doc.close()
    h = rgb.shape[0]
    crop = rgb[int(h * 0.68):int(h * 0.97)]
    Image.fromarray(crop).save(out_path, quality=94)


def flatten_audit(result):
    f = result.get("fields", {}) or {}
    v = result.get("visual", {}) or {}
    s = result.get("stamp_stage", {}) or {}
    row = {
        "attention_id": result.get("attention_id"),
        "source_pdf": result.get("source_pdf"),
        "consent_index": result.get("consent_index"),
        "page": result.get("page"),
        "procedure": result.get("procedure"),
        "template_type": result.get("template_type"),
    }

    field_geom_scores = (
        (result.get("field_presence_geometric_fallback", {}) or {}).get("scores", {}) or {}
    )

    for name in ("patient_name", "patient_rut", "procedure_date", "doctor_name", "doctor_rut"):
        x = f.get(name, {}) or {}
        row[f"{name}_status"] = x.get("status")
        row[f"{name}_confidence"] = x.get("confidence")
        row[f"{name}_source"] = x.get("source")
        row[f"{name}_geom_score"] = x.get(
            "geometric_presence_score", field_geom_scores.get(name)
        )
        row[f"{name}_manual_review"] = x.get("manual_review_required")
        row[f"{name}_reason"] = x.get("review_reason")

    cb = v.get("signer_selected_checkbox", {}) or {}
    cbg = cb.get("geometric_confirmation", {}) or {}
    ps = v.get("patient_signature", {}) or {}
    dsg = v.get("doctor_signature", {}) or {}
    data_stamp = v.get("data_stamp", {}) or {}

    row.update({
        "checkbox_status": cb.get("status"),
        "checkbox_option": cb.get("selected_option"),
        "checkbox_confidence": cb.get("confidence"),
        "checkbox_geom_score": cbg.get("presence_score"),
        "checkbox_geom_margin": cbg.get("option_margin"),

        "patient_signature_status": ps.get("status"),
        "patient_signature_confidence": ps.get("confidence"),
        "patient_signature_geom_score": (
            (ps.get("geometric_confirmation", {}) or {}).get("score")
        ),

        "doctor_signature_status": dsg.get("status"),
        "doctor_signature_confidence": dsg.get("confidence"),
        "doctor_signature_geom_score": (
            (dsg.get("geometric_confirmation", {}) or {}).get("score")
        ),

        "data_stamp_status": data_stamp.get("status"),
        "data_stamp_count": data_stamp.get("count"),
        "data_stamp_confidence": data_stamp.get("confidence"),

        "stamp_status": s.get("status"),
        "stamp_action": s.get("action"),
        "stamp_reason": s.get("reason"),

        "manual_review_required": result.get("manual_review_required"),
        "manual_review_reasons": " | ".join(
            f"{x.get('component')}:{x.get('reason')}"
            for x in (result.get("manual_review_items", []) or [])
        ),
        "output_pdf": result.get("output_pdf"),
        "final_pdf": result.get("final_pdf"),
        "rectified_image": result.get("rectified_image"),
    })
    return row


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------



def operational_preflight(args):
    """Valida rutas críticas antes de cargar modelos o procesar PDFs."""
    checks = [
        ("input", Path(args.input)),
        ("field model", Path(args.field_model)),
        ("visual main model", Path(args.visual_main)),
        ("doctor signature model", Path(args.doctor_model)),
        ("template alta", Path(args.template_alta)),
        ("template colon", Path(args.template_colon)),
        ("Excel atenciones", Path(args.atenciones_excel)),
        ("stamps dir", Path(args.stamps_dir)),
        ("Tesseract", Path(args.tesseract)),
    ]

    missing = [(label, str(path)) for label, path in checks if not path.exists()]
    if missing:
        print("=" * 72)
        print("INDISA V9.1 - PREFLIGHT ERROR")
        print("=" * 72)
        for label, path in missing:
            print(f"- {label}: NO ENCONTRADO -> {path}")
        raise FileNotFoundError(
            "Faltan archivos/rutas necesarias. Corrige las rutas antes de ejecutar."
        )

    inp = Path(args.input)
    if inp.is_dir():
        pdf_count = len(list(inp.glob("*.pdf")))
        if pdf_count == 0:
            raise RuntimeError(
                f"La carpeta de entrada no contiene PDFs: {inp}"
            )
    elif inp.suffix.lower() != ".pdf":
        raise RuntimeError(
            f"La entrada debe ser un PDF o una carpeta con PDFs: {inp}"
        )

    print("=" * 72)
    print("INDISA V9.1 - PREFLIGHT OK")
    print("=" * 72)
    print(f"Entrada: {args.input}")
    print(f"Salida: {args.output}")
    print("Modelos/plantillas/Excel/timbres/Tesseract: OK")
    print("Configuración congelada: SÍ")
    print("GT en ejecución: NO")
    print()


def _component_state(result, component):
    if component in result.get("fields", {}):
        x = result["fields"][component]
        return x.get("status"), x.get("confidence")
    if component in result.get("visual", {}):
        x = result["visual"][component]
        return x.get("status"), x.get("confidence")
    if component == "stamp_stage":
        x = result.get("stamp_stage", {})
        return x.get("status"), None
    return None, None


def _alert_message(component, reason, status):
    labels = {
        "patient_name": "Nombre del paciente",
        "patient_rut": "RUT del paciente",
        "procedure_date": "Fecha",
        "doctor_name": "Nombre del médico",
        "doctor_rut": "RUT del médico",
        "signer_selected_checkbox": "Tipo de firmante",
        "patient_signature": "Firma paciente/apoderado",
        "doctor_signature": "Firma médico",
        "stamp_stage": "Aplicación del timbre",
        "pdf": "PDF",
        "document_layout": "Estructura del documento",
        "rectification": "Rectificación",
        "inference": "Inferencia",
    }
    label = labels.get(component, component)

    if reason == "low_confidence":
        return f"{label}: inferencia de baja confianza; requiere revisión."
    if reason == "missing_required_field":
        return f"{label}: no se detectó contenido en el campo requerido."
    if reason == "not_detected":
        return f"{label}: no fue detectado."
    if reason == "ambiguous_multiple_candidates":
        return f"{label}: detección ambigua; hay múltiples candidatos."
    if reason == "doctor_signature_not_safely_located":
        return "No se pudo localizar la firma médica con seguridad para posicionar el timbre."
    if reason == "prior_stamp_inconclusive":
        return "Se detectó un posible timbre previo, pero no pudo confirmarse con seguridad."
    if reason == "pdf_alignment_low_confidence":
        return "La alineación con el PDF original no fue suficientemente confiable."
    if reason == "mapped_target_outside_page":
        return "La zona calculada para el timbre quedó fuera del área válida de la página."
    if reason == "doctor_stamp_png_not_found":
        return "No se encontró la imagen de timbre correspondiente al médico."
    if reason == "doctor_stamp_png_unreadable":
        return "La imagen de timbre no pudo leerse."
    if reason == "no_safe_target_inside_calibrated_slot":
        return "El timbre no cabe de forma segura dentro del slot calibrado."
    if reason and reason.startswith("missing_signature"):
        return f"{label}: faltan firmas necesarias para validar el slot de timbre."
    return f"{label}: {reason or status or 'requiere revisión'}."


def build_alerts_from_results(results):
    alerts = []
    for r in results:
        source_pdf = r.get("source_pdf") or Path(str(r.get("source_pdf_path", ""))).name
        page = r.get("page")
        for item in r.get("manual_review_items", []) or []:
            component = item.get("component")
            reason = item.get("reason")
            status, confidence = _component_state(r, component)
            alerts.append({
                "source_pdf": source_pdf,
                "page": page,
                "component": component,
                "status": status,
                "reason": reason,
                "confidence": confidence,
                "severity": "WARNING",
                "message": _alert_message(component, reason, status),
                "final_pdf": r.get("final_pdf") or r.get("output_pdf"),
            })
    return alerts


ANNOTATION_COLORS = {
    "FILLED": (0.0, 0.65, 0.0),
    "EMPTY": (1.0, 0.0, 0.0),
    "UNCERTAIN": (1.0, 0.55, 0.0),
    "CHECKBOX": (0.0, 0.65, 0.85),
    "SIGNATURE": (0.85, 0.0, 0.85),
    "STAMP": (0.45, 0.1, 0.75),
    "TARGET": (0.0, 0.8, 0.0),
}


def _rect_px_to_norm(box_px, rectified_shape):
    h, w = rectified_shape[:2]
    x1, y1, x2, y2 = [float(v) for v in box_px]
    return (x1 / w, y1 / h, x2 / w, y2 / h)


def _norm_box_to_pdf_quad(box_norm, rectified_shape, raw_shape, H, page):
    q = target_quad(box_norm, rectified_shape)
    mapped = cv2.perspectiveTransform(q, H).reshape(-1, 2)

    rh, rw = raw_shape[:2]
    pr = page.rect
    pts = []
    for x, y in mapped:
        px = pr.x0 + (float(x) / rw) * pr.width
        py = pr.y0 + (float(y) / rh) * pr.height
        pts.append(fitz.Point(px, py))
    return pts


def _draw_quad_with_label(page, pts, label, color, width=1.2):
    if len(pts) != 4:
        return
    poly = [pts[0], pts[1], pts[2], pts[3], pts[0]]
    page.draw_polyline(poly, color=color, width=width, overlay=True)
    p = pts[0]
    x = max(page.rect.x0 + 2, p.x)
    y = max(page.rect.y0 + 8, p.y - 2)
    try:
        page.insert_text(
            fitz.Point(x, y),
            label[:55],
            fontsize=5.5,
            color=color,
            overlay=True,
        )
    except Exception:
        pass


def _best_candidate_box(component):
    cands = component.get("candidates", []) or []
    valid = []
    for c in cands:
        b = c.get("bbox", {}).get("xyxy_norm")
        if b and len(b) == 4:
            valid.append(c)
    if not valid:
        return None
    c = max(valid, key=lambda z: float(z.get("confidence", 0) or 0))
    return tuple(float(v) for v in c["bbox"]["xyxy_norm"]), float(c.get("confidence", 0) or 0)


def annotate_final_pdf_v1(clean_pdf_path, annotated_pdf_path, source_pdf_path, page_results, dpi):
    """
    final_pdfs/ = PDF de revisión con cajas.
    pdfs/ y manual_review/ = PDF limpio.
    """
    src_pdf = Path(source_pdf_path)
    clean_pdf = Path(clean_pdf_path)
    annotated_pdf = Path(annotated_pdf_path)

    doc = fitz.open(str(clean_pdf))

    for r in page_results:
        try:
            page_num = int(r.get("page", 1))
            page_idx = page_num - 1
            if page_idx < 0 or page_idx >= len(doc):
                continue

            rect_path = Path(r.get("rectified_image", ""))
            rectified = safe_imread(rect_path, cv2.IMREAD_COLOR)
            if rectified is None:
                raise RuntimeError("rectified_image_unreadable")

            raw = render_pdf_page_bgr(src_pdf, page_num, dpi)
            H, hm = estimate_pdf_homography(rectified, raw)
            if (
                H is None
                or hm.get("inliers", 0) < MIN_H_INLIERS
                or hm.get("inlier_ratio", 0) < MIN_H_RATIO
                or hm.get("median_reprojection_error_px", 999) > MAX_H_MEDIAN_ERROR_PX
            ):
                raise RuntimeError("annotation_homography_low_confidence")

            page = doc[page_idx]

            # Campos estructurados
            for name, field in r.get("fields", {}).items():
                roi_px = field.get("roi_px")
                if not roi_px or len(roi_px) != 4:
                    continue
                status = field.get("status", "UNCERTAIN")
                conf = field.get("confidence")
                box = _rect_px_to_norm(roi_px, rectified.shape)
                pts = _norm_box_to_pdf_quad(box, rectified.shape, raw.shape, H, page)
                conf_txt = "" if conf is None else f" {float(conf):.2f}"
                _draw_quad_with_label(
                    page,
                    pts,
                    f"{name}: {status}{conf_txt}",
                    ANNOTATION_COLORS.get(status, ANNOTATION_COLORS["UNCERTAIN"]),
                    width=1.15,
                )

            # Checkbox
            checkbox = r.get("visual", {}).get("signer_selected_checkbox", {})
            bb = _best_candidate_box(checkbox)
            if bb:
                box, conf = bb
                pts = _norm_box_to_pdf_quad(box, rectified.shape, raw.shape, H, page)
                option = checkbox.get("selected_option") or "checkbox"
                _draw_quad_with_label(
                    page, pts, f"checkbox: {option} {conf:.2f}",
                    ANNOTATION_COLORS["CHECKBOX"], width=1.3
                )

            # Firmas
            for key, short in (
                ("patient_signature", "firma paciente"),
                ("doctor_signature", "firma medico"),
            ):
                comp = r.get("visual", {}).get(key, {})
                bb = _best_candidate_box(comp)
                if bb:
                    box, conf = bb
                    pts = _norm_box_to_pdf_quad(box, rectified.shape, raw.shape, H, page)
                    _draw_quad_with_label(
                        page, pts, f"{short}: {conf:.2f}",
                        ANNOTATION_COLORS["SIGNATURE"], width=1.5
                    )

            # Data stamp detectado
            stamp_comp = r.get("visual", {}).get("data_stamp", {})
            bb = _best_candidate_box(stamp_comp)
            if bb:
                box, conf = bb
                pts = _norm_box_to_pdf_quad(box, rectified.shape, raw.shape, H, page)
                _draw_quad_with_label(
                    page, pts, f"data_stamp: {conf:.2f}",
                    ANNOTATION_COLORS["STAMP"], width=1.4
                )

            # Target timbre si existió
            target = r.get("stamp_stage", {}).get("target_box_norm")
            if target and len(target) == 4:
                target = tuple(float(x) for x in target)
                pts = _norm_box_to_pdf_quad(target, rectified.shape, raw.shape, H, page)
                _draw_quad_with_label(
                    page, pts, "TARGET TIMBRE",
                    ANNOTATION_COLORS["TARGET"], width=2.2
                )

            # Resumen de revisión
            alerts = r.get("manual_review_items", []) or []
            if alerts:
                y = page.rect.y0 + 10
                x = page.rect.x0 + 8
                page.insert_text(
                    fitz.Point(x, y), "REVISAR:",
                    fontsize=6.5, color=(1.0, 0.0, 0.0), overlay=True
                )
                for item in alerts[:5]:
                    y += 7
                    msg = f"- {item.get('component')}: {item.get('reason')}"
                    page.insert_text(
                        fitz.Point(x, y), msg[:90],
                        fontsize=5.3, color=(0.8, 0.0, 0.0), overlay=True
                    )

            r["annotation_homography"] = hm

        except Exception as exc:
            r["annotation_error"] = f"{type(exc).__name__}: {exc}"

    if annotated_pdf.exists():
        annotated_pdf.unlink()
    doc.save(str(annotated_pdf), garbage=4, deflate=True)
    doc.close()



OUTPUT_SCHEMA_VERSION = "1.0"


def _simple_confirmation_from_visual(component):
    g = component.get("geometric_confirmation", {}) or {}
    if not g.get("confirmed"):
        return None

    out = {"method": "geometric_confirmation"}
    score = g.get("score", g.get("presence_score"))
    if score is not None:
        out["score"] = score
    if g.get("threshold") is not None:
        out["threshold"] = g.get("threshold")
    if g.get("option_margin") is not None:
        out["margin"] = g.get("option_margin")
    return out


def _simple_field(result, name):
    field = (result.get("fields", {}) or {}).get(name, {}) or {}
    out = {
        "status": field.get("status"),
        "confidence": field.get("confidence"),
        "source": field.get("source"),
    }

    source = field.get("source")
    if source == "template_subtraction_confirmation":
        geom_scores = (
            (result.get("field_presence_geometric_fallback", {}) or {})
            .get("scores", {}) or {}
        )
        score = field.get("geometric_presence_score", geom_scores.get(name))
        out["confirmation"] = {
            "method": "template_subtraction",
            "score": score,
        }

    # Quitar claves nulas para mantener el JSON compacto.
    return {k: v for k, v in out.items() if v is not None}


def _simple_visual(component, option_key=None):
    component = component or {}
    out = {
        "status": component.get("status"),
        "confidence": component.get("confidence"),
    }
    if option_key:
        out["option"] = component.get(option_key)

    confirmation = _simple_confirmation_from_visual(component)
    if confirmation:
        out["confirmation"] = confirmation

    return {k: v for k, v in out.items() if v is not None}


def simplify_result_document(result):
    fields = result.get("fields", {}) or {}
    visual = result.get("visual", {}) or {}
    stage = result.get("stamp_stage", {}) or {}
    alignment = result.get("template_alignment", {}) or {}

    source_pdf = result.get("source_pdf")
    if not source_pdf and result.get("source_pdf_path"):
        source_pdf = Path(result["source_pdf_path"]).name

    simple = {
        "source_pdf": source_pdf,
        "page": result.get("page"),
        "attention_id": result.get("attention_id"),
        "procedure": result.get("procedure"),
        "template": {
            "type": result.get("template_type"),
            "recognized": bool(result.get("template_type")),
            "alignment_inliers": alignment.get("orb_inliers"),
        },
        "fields": {
            name: _simple_field(result, name)
            for name in (
                "patient_name",
                "patient_rut",
                "procedure_date",
                "doctor_name",
                "doctor_rut",
            )
            if name in fields
        },
        "checkbox": _simple_visual(
            visual.get("signer_selected_checkbox", {}),
            option_key="selected_option",
        ),
        "signatures": {
            "patient": _simple_visual(visual.get("patient_signature", {})),
            "doctor": _simple_visual(visual.get("doctor_signature", {})),
        },
        "stamp": {
            "status": stage.get("status"),
            "action": stage.get("action"),
            "reason": stage.get("reason"),
            "existing_stamp_detected": stage.get("status") in {
                "EXISTING_STAMP_DETECTED",
                "EXISTING_CORRECT_STAMP",
            },
        },
        "review": {
            "required": bool(result.get("manual_review_required")),
            "reasons": [
                {
                    "component": x.get("component"),
                    "reason": x.get("reason"),
                }
                for x in (result.get("manual_review_items", []) or [])
            ],
        },
        "final_pdf": result.get("final_pdf"),
    }

    # Limpiar opcionales de nivel superior.
    return {
        k: v
        for k, v in simple.items()
        if v is not None
    }


def build_simple_results(results, pipeline_version):
    return {
        "schema_version": OUTPUT_SCHEMA_VERSION,
        "pipeline_version": pipeline_version,
        "documents": [simplify_result_document(r) for r in results],
    }

def main():
    ap = argparse.ArgumentParser(
        description="INDISA V9.1 - pipeline integrado operacional congelado"
    )
    ap.add_argument("--version", action="version", version="INDISA Pipeline V9.1")
    ap.add_argument("--input", required=True, help="PDF o carpeta con PDFs nuevos")
    ap.add_argument("--output", default=r"resultados_indisa_v9_1", help="Carpeta de resultados")

    ap.add_argument("--field-model", default=r"models\field_presence_v811_best.pt")
    ap.add_argument("--visual-main", default=r"models\visual_main_v813.pt")
    ap.add_argument("--doctor-model", default=r"models\doctor_signature_fallback_v811.pt")
    ap.add_argument(
        "--template-alta",
        default=r"datos\Consentimientos\Plantillas base\Consentimiento_base_EndoscopíaDigestivaAlta_Esófago-gastr.pdf",
    )
    ap.add_argument(
        "--template-colon",
        default=r"datos\Consentimientos\Plantillas base\Consentimiento_base_Colonoscopía_EndoscopíaDigestivaBaja.pdf",
    )
    ap.add_argument("--atenciones-excel", default=r"datos\Excel atenciones ejemplo.xlsx")
    ap.add_argument("--stamps-dir", default=r"datos\timbres")
    ap.add_argument("--tesseract", default=r"D:\Program Files\Tesseract\tesseract.exe")
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--device", default="0")
    ap.add_argument(
        "--allow-fixed-safe-zone",
        action="store_true",
        help=(
            "Permite timbrar sin bbox de firma usando zona fija. "
            "DESACTIVADO por defecto porque todavía no fue validado en TEST."
        ),
    )
    ap.add_argument(
        "--no-pdf-export",
        action="store_true",
        help=(
            "No genera copias PDF en pdfs/, manual_review/ ni final_pdfs/. "
            "Ejecuta inferencia y genera JSON/CSV sin escribir PDFs."
        ),
    )
    ap.add_argument("--overwrite", action="store_true")
    args = ap.parse_args()

    operational_preflight(args)

    out = Path(args.output)
    if out.exists() and any(out.iterdir()):
        if not args.overwrite:
            raise SystemExit(
                f"La salida ya existe y no está vacía: {out}\n"
                "Usa otra carpeta o agrega --overwrite."
            )
        shutil.rmtree(out)

    pdfs_dir = out / "pdfs"
    review_dir = out / "manual_review"
    final_dir = out / "final_pdfs"
    previews_dir = out / "previews"
    rectified_dir = out / "work" / "rectified"

    for d in (previews_dir, rectified_dir):
        d.mkdir(parents=True, exist_ok=True)

    if not args.no_pdf_export:
        for d in (pdfs_dir, review_dir, final_dir):
            d.mkdir(parents=True, exist_ok=True)

    # Validar archivos esenciales.
    essential = [
        Path(args.field_model), Path(args.visual_main), Path(args.doctor_model),
        Path(args.template_alta), Path(args.template_colon),
        Path(args.atenciones_excel), Path(args.stamps_dir),
    ]
    missing = [str(p) for p in essential if not p.exists()]
    if missing:
        raise SystemExit("Faltan recursos:\n- " + "\n- ".join(missing))

    pytesseract.pytesseract.tesseract_cmd = args.tesseract

    input_pdfs = collect_input_pdfs(args.input)
    if not input_pdfs:
        raise SystemExit("No encontré PDFs en --input.")

    worklist = load_worklist(Path(args.atenciones_excel))
    stamps = load_stamp_files(Path(args.stamps_dir))

    # Plantillas a la misma resolución usada para rectificación.
    tpl_alta = render_pdf_page_bgr(Path(args.template_alta), 1, args.dpi)
    tpl_colon = render_pdf_page_bgr(Path(args.template_colon), 1, args.dpi)
    templates = {"endoscopia_alta": tpl_alta, "colonoscopia": tpl_colon}

    print("Cargando modelos...")
    field_model = YOLO(args.field_model)
    visual_main = YOLO(args.visual_main)
    doctor_model = YOLO(args.doctor_model)
    print("Modelos cargados.")
    print("PDFs de entrada:", len(input_pdfs))
    print("Timbres disponibles:", len(stamps))
    print()

    all_results = []
    pdf_contexts = []
    fatal_pdf_errors = []

    for pdf_idx, pdf_path in enumerate(input_pdfs, start=1):
        aid = attention_id_from_name(pdf_path.name)
        print(f"[{pdf_idx}/{len(input_pdfs)}] {pdf_path.name} | ID={aid or '?'}")

        pdf_ctx = {
            "source_pdf": pdf_path,
            "attention_id": aid,
            "results": [],
            "jobs": [],
            "fatal_error": None,
        }

        if not aid or aid not in worklist:
            pdf_ctx["fatal_error"] = "attention_id_not_found_in_excel"
            fatal_pdf_errors.append({
                "source_pdf": str(pdf_path),
                "attention_id": aid,
                "reason": pdf_ctx["fatal_error"],
            })
            pdf_contexts.append(pdf_ctx)
            print("  -> REVIEW: ID no encontrado en Excel")
            continue

        try:
            doc_probe = fitz.open(str(pdf_path))
            page_count = len(doc_probe)
            doc_probe.close()
        except Exception as exc:
            pdf_ctx["fatal_error"] = f"pdf_open_failed:{type(exc).__name__}:{exc}"
            fatal_pdf_errors.append({
                "source_pdf": str(pdf_path),
                "attention_id": aid,
                "reason": pdf_ctx["fatal_error"],
            })
            pdf_contexts.append(pdf_ctx)
            print("  -> REVIEW: no se pudo abrir PDF")
            continue

        for ci, excel_rec in enumerate(worklist[aid], start=1):
            page_num = 1 + (ci - 1) * 3
            template_type = infer_template_type(excel_rec["procedure"])
            template_bgr = templates[template_type]

            base = {
                "schema_version": VERSION,
                "processed_at": now_iso(),
                "attention_id": aid,
                "source_pdf": pdf_path.name,
                "source_pdf_path": str(pdf_path),
                "consent_index": ci,
                "page": page_num,
                "procedure": excel_rec["procedure"],
                "template_type": template_type,
                "excel_reference": {
                    "doctor_rut": excel_rec["doctor_rut"],
                    "doctor_name": excel_rec["doctor_name"],
                    "patient_rut": excel_rec["patient_rut"],
                    "patient_name": excel_rec["patient_name"],
                },
            }

            if page_num > page_count:
                r = deepcopy(base)
                r.update({
                    "fields": {}, "visual": {},
                    "rectification": {"status": "FAILED", "reason": "expected_page_not_found"},
                    "stamp_stage": {
                        "status": "MANUAL_REVIEW",
                        "action": "NONE",
                        "reason": "expected_page_not_found",
                        "manual_review_required": True,
                        "review_reason": "expected_page_not_found",
                    },
                    "manual_review_required": True,
                    "manual_review_items": [{
                        "component": "document_layout",
                        "reason": "expected_page_not_found",
                    }],
                })
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW expected_page_not_found")
                continue

            try:
                raw_scan = render_pdf_page_bgr(pdf_path, page_num, args.dpi)
                H_orb, orb_matches, orb_inliers = orb_register(template_bgr, raw_scan)
                rectified = rectify_scan(raw_scan, H_orb, template_bgr.shape)
            except Exception as exc:
                r = deepcopy(base)
                r.update({
                    "fields": {}, "visual": {},
                    "rectification": {
                        "status": "FAILED",
                        "reason": f"{type(exc).__name__}: {exc}",
                    },
                    "stamp_stage": {
                        "status": "MANUAL_REVIEW",
                        "action": "NONE",
                        "reason": "rectification_failed",
                        "manual_review_required": True,
                        "review_reason": "rectification_failed",
                    },
                    "manual_review_required": True,
                    "manual_review_items": [{
                        "component": "rectification", "reason": "rectification_failed"
                    }],
                })
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW rectification_failed")
                continue

            stem = f"{aid}_c{ci}_{template_type}_p{page_num}"
            rect_path = rectified_dir / f"{stem}.jpg"
            safe_imwrite(rect_path, rectified, 95)

            try:
                r = analyze_rectified(
                    rectified,
                    template_bgr,
                    template_type,
                    field_model,
                    visual_main,
                    doctor_model,
                    args.device,
                )
            except Exception as exc:
                r = deepcopy(base)
                r.update({
                    "fields": {}, "visual": {},
                    "rectification": {
                        "status": "OK",
                        "orb_matches": orb_matches,
                        "orb_inliers": orb_inliers,
                    },
                    "rectified_image": str(rect_path),
                    "stamp_stage": {
                        "status": "MANUAL_REVIEW",
                        "action": "NONE",
                        "reason": "inference_failed",
                        "manual_review_required": True,
                        "review_reason": "inference_failed",
                    },
                    "manual_review_required": True,
                    "manual_review_items": [{
                        "component": "inference",
                        "reason": f"{type(exc).__name__}: {exc}",
                    }],
                })
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW inference_failed")
                continue

            r.update(base)
            r["rectification"] = {
                "status": "OK",
                "orb_matches": orb_matches,
                "orb_inliers": orb_inliers,
            }
            r["rectified_image"] = str(rect_path)

            # Anti-duplicado.
            anti = inspect_existing_doctor_stamp(
                rectified, r, excel_rec["doctor_rut"]
            )

            if anti["confirmed"]:
                r["stamp_stage"] = {
                    "status": "EXISTING_CORRECT_STAMP",
                    "action": "NO_INSERT",
                    "reason": "expected_doctor_rut_confirmed_in_existing_lower_stamp",
                    "manual_review_required": False,
                    "review_reason": None,
                    "anti_duplicate": anti,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(
                    f"  c{ci} p{page_num}: existing stamp | "
                    f"review={r['manual_review_required']}"
                )
                continue

            # Un sello inferior inconcluso NO bloquea, pero exige revisión.
            prior_inconclusive = anti["status"] == "INCONCLUSIVE_CANDIDATE"

            stamp_path = stamps.get(excel_rec["doctor_rut"])
            if not stamp_path:
                r["stamp_stage"] = {
                    "status": "MANUAL_REVIEW",
                    "action": "NONE",
                    "reason": "doctor_stamp_png_not_found",
                    "manual_review_required": True,
                    "review_reason": "doctor_stamp_png_not_found",
                    "anti_duplicate": anti,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW no stamp PNG")
                continue

            stamp_gray = normalize_stamp_gray(stamp_path)
            if stamp_gray is None:
                r["stamp_stage"] = {
                    "status": "MANUAL_REVIEW",
                    "action": "NONE",
                    "reason": "doctor_stamp_png_unreadable",
                    "manual_review_required": True,
                    "review_reason": "doctor_stamp_png_unreadable",
                    "anti_duplicate": anti,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW unreadable stamp PNG")
                continue

            sig_box = best_signature_bbox_norm(r)
            target, target_reason, placement_mode = propose_stamp_target(
                rectified.shape,
                stamp_gray,
                template_type,
                signature_box=sig_box,
                allow_fixed=args.allow_fixed_safe_zone,
            )

            if target is None:
                r["stamp_stage"] = {
                    "status": "MANUAL_REVIEW",
                    "action": "NONE",
                    "reason": target_reason,
                    "manual_review_required": True,
                    "review_reason": target_reason,
                    "anti_duplicate": anti,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW {target_reason}")
                continue

            # Homografía rectificada -> render real del PDF.
            try:
                H_pdf, hm = estimate_pdf_homography(rectified, raw_scan)
            except Exception as exc:
                H_pdf, hm = None, {"reason": f"{type(exc).__name__}: {exc}"}

            if (
                H_pdf is None
                or hm.get("inliers", 0) < MIN_H_INLIERS
                or hm.get("inlier_ratio", 0) < MIN_H_RATIO
                or hm.get("median_reprojection_error_px", 999) > MAX_H_MEDIAN_ERROR_PX
            ):
                r["stamp_stage"] = {
                    "status": "MANUAL_REVIEW",
                    "action": "NONE",
                    "reason": "pdf_alignment_low_confidence",
                    "manual_review_required": True,
                    "review_reason": "pdf_alignment_low_confidence",
                    "anti_duplicate": anti,
                    "placement_mode": placement_mode,
                    "target_box_norm": [round(float(x), 6) for x in target],
                    "homography": hm,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW pdf alignment")
                continue

            quad = cv2.perspectiveTransform(
                target_quad(target, rectified.shape), H_pdf
            ).reshape(-1, 2)
            rh, rw = raw_scan.shape[:2]
            if (
                np.any(quad[:, 0] < -5)
                or np.any(quad[:, 0] > rw + 5)
                or np.any(quad[:, 1] < -5)
                or np.any(quad[:, 1] > rh + 5)
            ):
                r["stamp_stage"] = {
                    "status": "MANUAL_REVIEW",
                    "action": "NONE",
                    "reason": "mapped_target_outside_page",
                    "manual_review_required": True,
                    "review_reason": "mapped_target_outside_page",
                    "anti_duplicate": anti,
                    "placement_mode": placement_mode,
                    "target_box_norm": [round(float(x), 6) for x in target],
                    "homography": hm,
                }
                recalc_manual_review(r)
                all_results.append(r)
                pdf_ctx["results"].append(r)
                print(f"  c{ci} p{page_num}: REVIEW mapped target outside page")
                continue

            overlay_png = build_page_overlay_png(raw_scan.shape, stamp_gray, quad)

            r["stamp_stage"] = {
                "status": "READY_TO_APPLY",
                "action": "INSERT",
                "reason": "ok",
                "manual_review_required": prior_inconclusive,
                "review_reason": "prior_stamp_inconclusive" if prior_inconclusive else None,
                "anti_duplicate": anti,
                "stamp_png": str(stamp_path),
                "placement_mode": placement_mode,
                "target_box_norm": [round(float(x), 6) for x in target],
                "mapped_quad_render_px": [
                    [round(float(x), 1), round(float(y), 1)] for x, y in quad
                ],
                "homography": hm,
            }
            recalc_manual_review(r)

            job = {
                "result": r,
                "page_idx": page_num - 1,
                "overlay_png": overlay_png,
            }
            pdf_ctx["jobs"].append(job)
            all_results.append(r)
            pdf_ctx["results"].append(r)
            print(
                f"  c{ci} p{page_num}: READY | "
                f"review={r['manual_review_required']} | placement={placement_mode}"
            )

        pdf_contexts.append(pdf_ctx)
        print()

    # ------------------------------------------------------------------
    # Escribir COPIAS PDF.
    # ------------------------------------------------------------------
    for ctx in pdf_contexts:
        src = ctx["source_pdf"]
        aid = ctx["attention_id"]

        # Modo rápido/API: conservar inferencias pero no escribir/copiar PDFs.
        if args.no_pdf_export:
            for r in ctx["results"]:
                r["output_pdf"] = None
                r["final_pdf"] = None
                r["final_pdf_mode"] = "disabled"
            continue

        if ctx["fatal_error"]:
            dst = review_dir / src.name
            shutil.copy2(src, dst)
            final_dst = final_dir / src.name
            shutil.copy2(dst, final_dst)
            continue

        # Un PDF va a manual_review si CUALQUIER consentimiento requiere revisión.
        pdf_needs_review = any(
            bool(r.get("manual_review_required")) for r in ctx["results"]
        )
        dest_dir = review_dir if pdf_needs_review else pdfs_dir
        dst = dest_dir / src.name

        try:
            if ctx["jobs"]:
                doc = fitz.open(str(src))
                for job in ctx["jobs"]:
                    page = doc[job["page_idx"]]
                    page.insert_image(
                        page.rect,
                        stream=job["overlay_png"],
                        overlay=True,
                        keep_proportion=False,
                    )
                    stage = job["result"]["stamp_stage"]
                    stage["status"] = "APPLIED"
                    stage["action"] = "INSERTED"
                doc.save(str(dst), garbage=4, deflate=True)
                doc.close()
            else:
                shutil.copy2(src, dst)

            final_dst = final_dir / src.name
            annotate_final_pdf_v1(
                dst,
                final_dst,
                src,
                ctx["results"],
                args.dpi,
            )

            for r in ctx["results"]:
                r["output_pdf"] = str(dst)          # limpio
                r["final_pdf"] = str(final_dst)     # anotado
                r["final_pdf_mode"] = "annotated_review_pdf"

            # Preview autoritativo del PDF guardado por cada consentimiento.
            for r in ctx["results"]:
                page_idx = int(r.get("page", 1)) - 1
                try:
                    preview_name = (
                        f"{aid}_c{r.get('consent_index')}_p{r.get('page')}_after.jpg"
                    )
                    save_preview_from_pdf(final_dst, page_idx, previews_dir / preview_name)
                    r["preview"] = str(previews_dir / preview_name)
                except Exception as exc:
                    r["preview_error"] = f"{type(exc).__name__}: {exc}"

        except Exception as exc:
            # Si la escritura falla, conservar original en manual_review.
            fail_dst = review_dir / src.name
            if not fail_dst.exists():
                shutil.copy2(src, fail_dst)
            final_dst = final_dir / src.name
            shutil.copy2(fail_dst, final_dst)
            for r in ctx["results"]:
                r["output_pdf"] = str(fail_dst)
                r["final_pdf"] = str(final_dst)
                r.setdefault("stamp_stage", {})["write_error"] = f"{type(exc).__name__}: {exc}"
                r["stamp_stage"]["manual_review_required"] = True
                r["stamp_stage"]["review_reason"] = "pdf_write_failed"
                recalc_manual_review(r)

    # Fatal PDFs sin consentimiento procesable: agregarlos al resultado global.
    for e in fatal_pdf_errors:
        all_results.append({
            "schema_version": VERSION,
            "record_type": "pdf_error",
            "processed_at": now_iso(),
            "attention_id": e["attention_id"],
            "source_pdf_path": e["source_pdf"],
            "manual_review_required": True,
            "manual_review_items": [{
                "component": "pdf", "reason": e["reason"]
            }],
            "stamp_stage": {
                "status": "MANUAL_REVIEW",
                "action": "NONE",
                "reason": e["reason"],
            },
        })

    # ------------------------------------------------------------------
    # Resultados y auditoría.
    # ------------------------------------------------------------------
    run_metadata = {
        "pipeline_version": VERSION,
        "started_or_finished_at": now_iso(),
        "input": str(Path(args.input)),
        "output": str(out),
        "gt_used": False,
        "thresholds_recalibrated": False,
        "originals_modified": False,
        "allow_fixed_safe_zone": bool(args.allow_fixed_safe_zone),
        "models": {
            "field_presence": args.field_model,
            "visual_main": args.visual_main,
            "doctor_signature": args.doctor_model,
        },
    }

    debug_results_payload = {
        "run": run_metadata,
        "documents": all_results,
    }

    # Fuente técnica completa para trazabilidad / depuración.
    (out / "results_debug.json").write_text(
        json.dumps(debug_results_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    # Mismo contrato simplificado que V2.
    simple_results_payload = build_simple_results(all_results, VERSION)
    (out / "results.json").write_text(
        json.dumps(simple_results_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    alerts = build_alerts_from_results(all_results)
    alerts_payload = {
        "schema_version": OUTPUT_SCHEMA_VERSION,
        "pipeline_version": VERSION,
        "alert_count": len(alerts),
        "alerts": alerts,
    }
    (out / "alerts.json").write_text(
        json.dumps(alerts_payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    audit_rows = [flatten_audit(r) for r in all_results if r.get("record_type") != "pdf_error"]
    audit_path = out / "audit_log.csv"
    if audit_rows:
        with audit_path.open("w", newline="", encoding="utf-8-sig") as f:
            wr = csv.DictWriter(f, fieldnames=list(audit_rows[0].keys()))
            wr.writeheader()
            wr.writerows(audit_rows)
    else:
        audit_path.write_text("", encoding="utf-8")

    processed_docs = [r for r in all_results if r.get("record_type") != "pdf_error"]
    summary = {
        "schema_version": OUTPUT_SCHEMA_VERSION,
        "pipeline_version": VERSION,
        "pdfs_input": len(input_pdfs),
        "consents_processed": len(processed_docs),
        "manual_review_consents": sum(bool(r.get("manual_review_required")) for r in processed_docs),
        "stamp_applied": sum(r.get("stamp_stage", {}).get("status") == "APPLIED" for r in processed_docs),
        "existing_correct_stamp": sum(
            r.get("stamp_stage", {}).get("status") == "EXISTING_CORRECT_STAMP"
            for r in processed_docs
        ),
        "stamp_manual_review_or_blocked": sum(
            r.get("stamp_stage", {}).get("status") == "MANUAL_REVIEW"
            for r in processed_docs
        ),
        "pdf_level_errors": len(fatal_pdf_errors),
        "clean_output_pdfs": 0 if args.no_pdf_export else len(list(pdfs_dir.glob("*.pdf"))),
        "manual_review_pdfs": 0 if args.no_pdf_export else len(list(review_dir.glob("*.pdf"))),
        "pdf_export_enabled": not args.no_pdf_export,
        "originals_modified": 0,
        "gt_used": False,
        "thresholds_recalibrated": False,
        "alerts_total": len(alerts),
    }
    (out / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("=" * 72)
    print("INDISA V1.1 / V9.1 PRODUCTION ANNOTATED - RESUMEN")
    print("=" * 72)
    print("PDFs input:", summary["pdfs_input"])
    print("Consentimientos procesados:", summary["consents_processed"])
    print("Timbres aplicados:", summary["stamp_applied"])
    print("Timbre correcto ya existente:", summary["existing_correct_stamp"])
    print("Consentimientos con manual review:", summary["manual_review_consents"])
    print("PDFs clean output:", summary["clean_output_pdfs"])
    print("PDFs manual_review:", summary["manual_review_pdfs"])
    print("Errores PDF-level:", summary["pdf_level_errors"])
    print("Originales modificados: 0")
    print("GT utilizado: NO")
    print("Thresholds recalibrados: NO")
    print()
    print("Resultados simples:", out / "results.json")
    print("Resultados debug:", out / "results_debug.json")
    print("Auditoría:", audit_path)
    print("Resumen:", out / "summary.json")
    print("PDFs limpios:", pdfs_dir)
    print("PDFs a revisión:", review_dir)
    print("Previews:", previews_dir)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("Interrumpido por usuario.", file=sys.stderr)
        raise
    except Exception as exc:
        print(f"ERROR FATAL: {type(exc).__name__}: {exc}", file=sys.stderr)
        traceback.print_exc()
        raise
